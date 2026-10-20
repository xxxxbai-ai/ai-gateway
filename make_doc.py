# -*- coding: utf-8 -*-
"""生成《AI 网关问题分类器 实现原理详解.docx》"""
import docx
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MONO_ASCII = "Consolas"
CN_BODY = "微软雅黑"
CN_EAST = "微软雅黑"

doc = Document()

# ---------- 工具函数 ----------
def set_run(run, ascii_font=CN_BODY, east_font=CN_EAST, size=None, bold=None, color=None):
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)
    rFonts.set(qn('w:eastAsia'), east_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=22, bold=True)
    p.paragraph_format.space_after = Pt(6)
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=11, color=(0x60, 0x60, 0x60))
    p.paragraph_format.space_after = Pt(14)
    return p

def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=16, bold=True, color=(0x1F, 0x3B, 0x63))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=13, bold=True, color=(0x2E, 0x54, 0x86))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def body(text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    set_run(r, size=10.5)
    p.paragraph_format.space_after = Pt(2)
    return p

def shade(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def code_block(text, fill="F5F5F5"):
    lines = text.strip('\n').split('\n')
    for ln in lines:
        p = doc.add_paragraph()
        r = p.add_run(ln if ln else ' ')
        set_run(r, ascii_font=MONO_ASCII, east_font=CN_EAST, size=9)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Pt(8)
        shade(p, fill)

def inline_code(para_text_parts):
    """parts: list of (text, is_code)"""
    p = doc.add_paragraph()
    for text, is_code in para_text_parts:
        r = p.add_run(text)
        if is_code:
            set_run(r, ascii_font=MONO_ASCII, east_font=CN_EAST, size=10)
        else:
            set_run(r, size=10.5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p

def add_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        r = hdr[i].paragraphs[0].add_run(h)
        set_run(r, size=10, bold=True)
        shade(hdr[i].paragraphs[0], "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(v)
            set_run(r, size=9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# ================= 封面标题 =================
title("AI 网关问题分类器 · 实现原理详解")
subtitle("规则引擎式意图分类：代码 / 数学计算 / 翻译 / 写作 / 通用 五分类")

# ================= 一、概述 =================
h1("一、概述")
body("本分类器是 AI 网关（AI Gateway）的一道前置路由组件。网关在收到用户问题后，需要决定把问题转交给哪一个专业模型处理："
     "数学计算模型、代码模型、翻译模型、写作模型，还是通用对话模型。分类器的职责就是在正式调用大模型之前，用极低的成本判断出问题的意图类别。")
body("设计目标只有两个：快、省。它完全基于规则（关键词 + 正则表达式）打分，不调用任何大模型，单条分类耗时在毫秒级，"
     "也不需要 GPU。这使它适合作为网关的第一道「粗路由」，把昂贵的大模型算力留给真正需要它的问题。")
inline_code([("核心技术手段：", False), ("关键词字典 + 正则表达式 + 加权打分 + 最高分决策", True), ("。", False)])

# ================= 二、整体架构与执行流程 =================
h1("二、整体架构与执行流程")
body("程序由四部分组成：")
bullet("main() —— 入口，负责解析三种输入模式（--test 自测 / 命令行参数 / 标准输入）并输出结果。")
bullet("QueryClassifier —— 核心类，封装五类关键词字典、正则字典与 classify() 分类逻辑。")
bullet("辅助函数 —— asciiLower()（归一化）、countMatches()（正则计数）、acpToUtf8()（Windows 编码转换）、printResult()（输出）。")
bullet("selfTest() —— 内置 21 条用例的回归测试，保证改完规则后行为可校验。")

h2("2.1 执行流程")
code_block("""
用户问题 (std::string)
      │
      ▼
asciiLower()  归一化：仅 ASCII 字母转小写，中文 UTF-8 原样保留
      │
      ▼
score()  对 5 个类别分别打分
      │   代码=关键词分+正则分
      │   数学=关键词分+正则分
      │   翻译=关键词分+正则分
      │   写作=关键词分+正则分
      │   通用=关键词分+正则分+基线分(1)
      ▼
argmax 决策：取最高分，平局按枚举优先级(代码>数学>翻译>写作>通用)
      │
      ▼
confidence = 最高分 ÷ 五类总分
      │
      ▼
输出：[类别] 问题文本 + 各类评分明细 + 置信度
""")

body("整个流程是确定性的——同样的输入永远得到同样的输出，没有随机性，便于测试和调试。这一点对网关这种基础设施尤其重要。")

# ================= 三、核心数据结构 =================
h1("三、核心数据结构")
h2("3.1 类别枚举 Category")
body("枚举的取值顺序不是随意的，它同时定义了「平局时的优先级」：靠前者胜出。")
code_block("""
enum class Category { Code = 0, Math, Translation, Writing, General };
//  顺序即优先级：  代码 > 数学 > 翻译 > 写作 > 通用
""")
body("之所以把「代码」排第一，是因为代码类问题常常混入大量动作词（如「帮我写」「实现」「怎么写」），"
     "与写作类高度重叠。在两者同分时优先判为代码，更符合工程直觉。")

h2("3.2 分类结果 Result")
code_block("""
struct Result {
    Category category;    // 胜出类别
    double   confidence;  // 置信度 0~1
    int      scores[5];   // 五个类别的原始得分，索引 = Category 枚举值
};
""")
body("scores[5] 保留了每类的原始得分，不只是给出最终类别。这让结果可解释——网关或人类都能看到「为什么判成这一类」。")

h2("3.3 规则字典")
body("每个类别持有两份字典，一份关键词、一份正则，都是「模式 → 权重」的键值对：")
code_block("""
std::vector<std::pair<std::string, int>> kws_[5];  // 关键词 -> 权重
std::vector<std::pair<std::regex,  int>> res_[5];  // 正则   -> 权重
""")
body("关键词负责「语义」层面的信号（出现「报错」「翻译」「作文」这些词），正则负责「结构」层面的信号（出现 1+2、#include、{ } ; 这类模式）。"
     "两者互补，是这套规则系统能覆盖五类问题的关键。")

# ================= 四、文本归一化 =================
h1("四、文本归一化 asciiLower()")
body("分类前先把英文统一成小写，实现「大小写不敏感」匹配，让 'Python'、'PYTHON'、'python' 都能命中关键词 'python'。")
code_block("""
std::string asciiLower(const std::string& s) {
    std::string r = s;
    for (char& c : r)
        if ((unsigned char)c < 0x80) c = (char)std::tolower((unsigned char)c);
    return r;
}
""")
body("这里有一个容易被忽略但很重要的细节：只对 ASCII 字符（字节值 < 0x80）转小写。")
body("因为 UTF-8 编码下，一个汉字由 3 个字节组成，且每个字节的首位都是 1（字节值 ≥ 0x80）。"
     "如果把整个字符串按字节调用 tolower()，会破坏中文的多字节序列，导致乱码或误匹配。"
     "上面的判断 (unsigned char)c < 0x80 精确地把中文字节挡在外面，中文内容原样保留。")

# ================= 五、打分机制 =================
h1("五、打分机制 score()")
body("打分是分类器的核心。对某个类别 q 的得分 = 关键词命中总分 + 正则命中总分：")
code_block("""
int score(Category c, const std::string& q) const {
    int idx = (int)c, s = 0;
    for (auto& kw : kws_[idx])                     // 关键词
        if (q.find(kw.first) != std::string::npos) // 子串包含即命中
            s += kw.second;
    for (auto& re : res_[idx])                     // 正则
        s += re.second * countMatches(re.first, q);
    return s;
}
""")

h2("5.1 关键词匹配：子串包含")
body("用 std::string::find() 做子串匹配。对中文这种「整词即子串」的语言，子串匹配天然安全——"
     "一个汉字的 UTF-8 字节序列不可能错误地横跨到相邻汉字中间，所以不会出现跨边界误匹配的问题。")

h2("5.2 正则匹配：计数加权")
body("与关键词的「命中一次加一次权重」不同，正则采用「命中次数 × 权重」的累加方式：")
code_block("""
static int countMatches(const std::regex& re, const std::string& s) {
    int n = 0;
    for (std::sregex_iterator it(s.begin(), s.end(), re), end; it != end; ++it) ++n;
    return n;
}
""")
body("为什么要计数？因为结构特征往往在一个句子里多次出现。例如 for (int i=0;i<n;i++) { } 里，"
     "关键字 for、int 各命中一次，花括号和分号命中 4 次——计数让这类「代码骨架」问题获得足够高的分数，"
     "而不只是象征性的 1 分。数学表达式的 1+2+3 同理，每个运算符都独立贡献分数。")

h2("5.3 权重设计")
body("权重取值集中在 1~4 之间，规则是「区分度越高的词权重越大」：")
add_table(
    ["权重", "含义", "举例"],
    [
        ["4", "极强信号，几乎唯一指向该类别", "翻译、translate、#include、帮我写一篇"],
        ["3", "强信号，非常典型", "报错、代码、爬虫、作文、方程、平方根、正则表达式"],
        ["2", "中等信号，常见但可能跨类", "函数、编译、百分比、邮件、总结"],
        ["1", "弱信号，单独出现不足以定类", "写、算、求、怎么、如何、介绍"],
    ],
    widths=[0.6, 2.6, 2.8],
)
body("弱信号词（权重 1）本身不足以决定类别，但多个弱信号叠加仍能形成判断；"
     "它们主要起「辅助确认」作用，真正的定类靠高权重词和正则。")

# ================= 六、分类决策 =================
h1("六、分类决策与兜底")
h2("6.1 最高分决策")
code_block("""
int best = 0, total = 0;
for (int i = 0; i < 5; ++i) {
    total += r.scores[i];
    if (r.scores[i] > r.scores[best]) best = i;  // 严格大于 => 平局取枚举顺序靠前者
}
r.category = (Category)best;
""")
body("决策非常简单：五类里分数最高者胜出。用「严格大于 >」而不是「大于等于 >=」遍历，"
     "意味着当两个类别同分时，枚举顺序靠前的类别获胜，即前面提到的优先级 代码 > 数学 > 翻译 > 写作 > 通用。")

h2("6.2 通用类的基线分兜底")
code_block("""
r.scores[(int)Category::General] = generalBase_ + score(Category::General, q);
// generalBase_ = 1
""")
body("「通用」类额外加了一个固定基线分 1。这是兜底机制的核心：当代码/数学/翻译/写作四类一个信号都没命中（全为 0 分）时，"
     "通用类至少拿到 1 分，从而自然胜出。没有这个基线分，一个完全匹配不到任何规则的问题会得到全零分，"
     "决策函数将无从下手。基线分保证「总能给出一个答案」，且答案合理——匹配不到任何专业特征的问题，本来就该交给通用模型。")

# ================= 七、置信度计算 =================
h1("七、置信度计算")
code_block("""
r.confidence = total > 0 ? (double)r.scores[best] / total : 0.0;
""")
body("置信度 = 胜出类别得分 ÷ 五类总分。它衡量的是这次判别的「干净程度」：")
bullet("置信度 = 1.0：只有胜出类得分，其余四类全 0。判别非常干净（如纯知识问答类问题）。")
bullet("置信度较低（如 0.5~0.6）：多个类别同时得分，信号有重叠。例如「帮我写一个 Python 爬虫」同时命中写作和代码，"
       "最终代码以微弱优势胜出，置信度 0.6，如实反映了这种模糊性。")
body("对网关而言，置信度是个有用的附加信息：低于阈值（例如 0.5）的问题可以降级为「通用」处理，或让大模型做二次确认。")

# ================= 八、Windows 编码处理 =================
h1("八、Windows 编码处理")
body("这段程序在 Windows 上跑，需要处理两个编码问题，否则中文会乱码或匹配失败。")
h2("8.1 控制台 UTF-8")
code_block("""
#ifdef _WIN32
    SetConsoleOutputCP(CP_UTF8);
    SetConsoleCP(CP_UTF8);
#endif
""")
body("程序内部统一使用 UTF-8 处理字符串，所以把控制台的输出/输入代码页都设成 UTF-8，保证打印和读入的中文正确。")
h2("8.2 命令行参数 GBK → UTF-8")
body("Windows 中文系统（GBK 区域）下，main 的 argv 参数由 C 运行库用系统 ANSI 代码页（GBK）解码，"
     "得到的不是 UTF-8。直接拿去和 UTF-8 关键词比对会匹配不上。于是用 Windows API 做一次转码：")
code_block("""
// GBK(CP_ACP) -> UTF-8
MultiByteToWideChar(CP_ACP, 0, in.c_str(), -1, ...)   // GBK -> UTF-16
WideCharToMultiByte(CP_UTF8, 0, w.c_str(), -1, ...)    // UTF-16 -> UTF-8
""")
body("经过这两步，命令行传入的中文问题（如 classifier.exe \"帮我写一封辞职邮件\"）就能被正确识别为「写作」。")

# ================= 九、五类判据详解 =================
h1("九、五个类别的具体判据")
add_table(
    ["类别", "关键词特征（语义）", "正则特征（结构）"],
    [
        ["代码",
         "代码、程序、编程、函数、报错、bug、调试、指针、多线程、爬虫、接口、API、数据库、算法、递归；\n"
         "语言名：python、c++、java、javascript、rust、golang、php、汇编…",
         "#include、import x、def f()、class/int/for/return 等关键字、{ } ; 括号分号"],
        ["数学计算",
         "计算、等于、方程、求解、平方、立方、次方、幂、根号、积分、导数、阶乘、质数、最大公约数、进制、百分比、换算…",
         "1+2、3x4、x=1、sin()、sqrt()、含运算符的整句算式"],
        ["翻译",
         "翻译、translate、翻译成、译成、翻成、中译英、英译中、什么意思、怎么读、发音、英文/日语/法语…",
         "translate ... to ..."],
        ["写作",
         "作文、文章、论文、报告、邮件、文案、演讲稿、总结、摘要、润色、改写、扩写、缩写、大纲、帮我写…",
         "（无强正则，主要靠文体名词定类）"],
        ["通用",
         "什么是、为什么、怎么、如何、介绍、解释、区别、推荐、原理、原因、怎么样…",
         "（兜底类别，加基线分）"],
    ],
    widths=[0.9, 2.7, 2.4],
)
body("三个容易混淆的边界，规则做了专门处理：")
bullet("「写代码」 vs 「写作文」：把「帮我写」「写」这类动作词降为低权重，靠宾语名词（代码/函数 vs 文章/邮件）最终定类。")
bullet("「计算」既可能是数学也可能是代码（如「用 C++ 计算」）：数学关键词与语言名关键词分别计分，由分数高低裁决。")
bullet("「这个单词怎么读」：同时含翻译信号（怎么读）和通用信号（怎么），翻译词权重更高，最终判翻译。")

# ================= 十、关键设计权衡 =================
h1("十、关键设计权衡")
h2("10.1 为什么用规则而不是大模型/机器学习")
add_table(
    ["方案", "优点", "缺点", "适用"],
    [
        ["规则分类（本方案）", "毫秒级、零成本、可解释、可回归测试", "覆盖有限、需人工维护字典", "网关第一道路由"],
        ["机器学习（朴素贝叶斯/TF-IDF）", "自动学习、泛化好", "需标注数据、冷启动难、可解释性差", "规则兜底后的补充"],
        ["大模型直接分类", "准确率最高、零维护", "慢、贵、可能不稳定", "高价值请求的二次确认"],
    ],
    widths=[1.8, 1.9, 1.9, 1.4],
)
h2("10.2 三条核心原则")
bullet("中性词降权：动作词（帮我写、实现）本身不表态，由它后面的名词表态。避免动作词误导分类。")
bullet("通用兜底：基线分保证永远有答案，且「匹配不到专业特征 → 通用」这一语义天然正确。")
bullet("正则补结构：关键词抓语义，正则抓结构。只靠关键词无法识别「1+2」这种没有语义词的纯算式；只靠正则无法识别「帮我写个爬虫」。两者缺一不可。")

# ================= 十一、完整源码 =================
h1("十一、完整源码（附录）")
body("完整实现见 D:\\query_classifier\\classifier.cpp。核心逻辑已在前面章节逐段讲解，以下是全文：", size=10.5)
try:
    with open(r"D:\query_classifier\classifier.cpp", "r", encoding="utf-8") as f:
        src = f.read()
    code_block(src)
except Exception as e:
    body("（读取源码失败：" + str(e) + "）")

# ================= 十二、扩展方向 =================
h1("十二、扩展方向")
bullet("配置外置：把关键词/权重/正则移到 JSON 或 YAML 配置，改规则不用重编译。")
bullet("统计层补充：对规则低置信度的问题，再用朴素贝叶斯/TF-IDF 做第二层判断。")
bullet("TOP-N 输出：不只给最优类别，输出带分数排序的前 N 个候选，供网关做加权路由（多模型混合）。")
bullet("大模型终审：高价值请求交给大模型二次确认，形成「规则粗分 → 模型精分」的两级架构。")

# ---------- 页边距与保存 ----------
for sec in doc.sections:
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)

out = r"D:\query_classifier\AI网关问题分类器_实现原理详解.docx"
doc.save(out)
print("已生成:", out)
